# -*- coding: utf-8 -*-
"""交付④ Word 可跳转年限 docx
结构:
- 封面/说明页
- 目录页: 五层(T0..T3)书签超链接 + 层内子时段书签超链接(点击跳转)
- 正文: 每层 Heading1, 每个子时段 Heading2 + 书签 bookmark, 子时段下列代表事件(标题+时间+锚点+一句话)
用 python-docx + 原生 XML 添加 bookmark 与内部超链接(anchor)。
"""
import os, json, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"D:/DeepSeek Harness专用！危险！！！！！/迷深清洗工作2/年表输出"
DELIVER = os.path.join(BASE, "交付四件套")
fl = json.load(open(os.path.join(BASE, "final_output", "full_event_list.json"), encoding="utf-8"))
agg = json.load(open(os.path.join(DELIVER, "_agg.json"), encoding="utf-8"))
by_id = {e["event_id"]: e for e in fl}

LAYER_META = {
    "T0": ("段0 · 原世界(现代日本)", "58 事件 ｜ 原生家庭·英才教育·妹妹病倒·病床前救赎", 1),
    "T1": ("段1 · 千年前·战争期(T-1000~T-995)", "123 事件 ｜ 召唤·圣人救世·阳滝异化·始祖复仇", 1),
    "T2": ("段2 · 千年前·建造期(T-995~T-990)", "119 事件 ｜ 十理就位·诺斯菲诞生·形式婚姻·涡波休眠", 1),
    "T4": ("段3 · 千年间(T-990~T-0)", "45 事件 ｜ 拉丝缇娅拉培育·守护者就位·攻略至23层", 1),
    "T3": ("段4 · 千年后主线(T-0~终章)", "1223 事件 ｜ 一层苏醒·逐层攻略·守护者战·终局", 1),
}

def add_bookmark(paragraph, bookmark_id, bookmark_name):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    paragraph._p.append(start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(end)

def add_internal_hyperlink(paragraph, bookmark_name, text, font_size=None, bold=False, color="0563C1"):
    part = paragraph.part
    r_id = part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=False)
    # 内部 anchor 超链接需在 rels 中为 internal? python-docx 方式: 直接构造 w:hyperlink w:anchor
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark_name)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if font_size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(font_size * 2))); rPr.append(sz)
    if bold:
        b = OxmlElement("w:b"); rPr.append(b)
    col = OxmlElement("w:color"); col.set(qn("w:val"), color); rPr.append(col)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t"); t.text = text; new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

doc = Document()
# 默认字体(中文字体)
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# 封面
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("《以异世界迷宫最深处为目标》(web版)"); r.font.size = Pt(20); r.bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("事件年表 · 按年限可跳转版"); r.font.size = Pt(16); r.bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("千年双时间线(时间轴序): 原世界 → 千年前·战争期 → 千年前·建造期 → 千年间 → 千年后主线(层编号为工作号,时间序见 time_seq)")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("数据源: final_output/full_event_list.json(1568 条) ｜ 生成 2026-09-01")
doc.add_page_break()

# 目录页
doc.add_heading("目录(点击跳转)", level=1)
bm_id = 1000
# 记录 (label -> bookmarkName, level)
toc_items = []  # (level, label, bookmark)

bookmarks = {}
for L in ("T0", "T1", "T2", "T4", "T3"):
    name, meta, _ = LAYER_META[L]
    bk = f"SEC_{L}"
    bookmarks[f"{L}_sec"] = bk
    toc_items.append((1, name, bk))
    for seg in agg["layers"][L]:
        segbk = f"SEG_{L}_{seg['seg_name'][:12]}"
        if L == "T3":
            toc_items.append((2, f"{seg['seg_name']} · {seg['seg_desc']}({len(seg['event_ids'])}条)", segbk))
        else:
            toc_items.append((2, f"{seg['seg_name']} ({len(seg['event_ids'])}条)", segbk))

# 写目录链接
for lvl, label, bk in toc_items:
    p = doc.add_paragraph()
    if lvl == 2:
        p.paragraph_format.left_indent = Cm(0.8)
    add_internal_hyperlink(p, bk, label, font_size=(11 if lvl == 1 else 10), bold=(lvl == 1))

doc.add_page_break()

# 正文
for L in ("T0", "T1", "T2", "T4", "T3"):
    name, meta, _ = LAYER_META[L]
    bk = bookmarks[f"{L}_sec"]
    h = doc.add_heading(name, level=1)
    add_bookmark(h, len(bookmarks) + len(toc_items) + 1, bk)
    doc.add_paragraph(meta)
    for seg in agg["layers"][L]:
        segbk = f"SEG_{L}_{seg['seg_name'][:12]}"
        h2 = doc.add_heading(f"{seg['seg_name']} · {seg['seg_desc'][:40]}{'…' if len(seg['seg_desc'])>40 else ''}({len(seg['event_ids'])}条)", level=2)
        add_bookmark(h2, len(bookmarks) + len(toc_items) + 2, segbk)
        # 代表事件(≤6): 标题+时间+锚点+首句
        ids = seg["event_ids"]
        top = []
        for i in ids:
            e = by_id[i]
            top.append(e)
        top.sort(key=lambda e: (0 if e.get("confidence") == "原文明确" else 1, e["event_id"]))
        for e in top[:6]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(f"{e['event_id']} ｜ {e['title']}")
            r.bold = True
            p2 = doc.add_paragraph()
            p2.paragraph_format.left_indent = Cm(0.6)
            d = (e.get("description", "") or "").replace("\n", " ")
            m = re.split(r"(?<=[。！？!?])", d)
            sent = m[0][:150] if m else d[:150]
            p2.add_run(f"时间: {e.get('event_time','')} ｜ 类型: {e.get('event_type','')} ｜ 锚点: {e.get('story_anchor','')}")
            p2.add_run("\n" + sent)
        if len(ids) > 6:
            doc.add_paragraph(f"— 其余 {len(ids)-6} 条详见 01_最终年表_全量.md —")
        # 主线小目录: T3 章内按话组分节标题(方案第5项; 节级细读指向 01)
        if L == "T3":
            from _交付_节映射 import get_secs_by_chapter
            secs = get_secs_by_chapter(seg["seg_name"])
            if len(secs) > 1:
                psec = doc.add_paragraph()
                rsec = psec.add_run(f"本章分 {len(secs)} 节(话组):")
                rsec.bold = True
                for s in secs:
                    nm = s["sec_name"].split("·", 1)[1] if "·" in s["sec_name"] else s["sec_name"]
                    pn = doc.add_paragraph()
                    pn.paragraph_format.left_indent = Cm(0.6)
                    rn = pn.add_run(f"▸ {nm}({len(s['event_ids'])}条)")
                    rn.font.color.rgb = RGBColor.from_string("4A3A8C")
    doc.add_page_break()

outp = os.path.join(DELIVER, "04_事件年表_按年限跳转_v5.docx")
doc.save(outp)
print("④ saved:", outp)
